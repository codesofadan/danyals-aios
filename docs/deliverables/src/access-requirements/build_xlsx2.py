"""Build the REVISED access-requirements spreadsheet. See data2.py for what changed."""
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import data2 as D

T = D.TOTALS
OUT = pathlib.Path(
    r"C:/Users/adan/Desktop/danyals-aios/docs/deliverables/danyal-aios-access-requirements.xlsx"
)
OUT.parent.mkdir(parents=True, exist_ok=True)

MAROON = "FF6E1423"
BLUSH = "FFF8ECEE"
LINE = "FFE8D2D7"
INK = "FF241015"
H = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
B = Font(name="Calibri", size=10.5, color=INK)
TITLE = Font(name="Calibri", size=15, bold=True, color=MAROON)
hfill = PatternFill("solid", fgColor=MAROON)
zfill = PatternFill("solid", fgColor=BLUSH)
thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wb.remove(wb.active)


def sheet(name, title, subtitle, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="FF6B5158")
    ws.append([])
    ws.append(headers)
    hr = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = H
        cell.fill = hfill
        cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
        rr = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = B
            cell.border = box
            cell.alignment = wrap
            if rr % 2 == 0:
                cell.fill = zfill
    for i, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.row_dimensions[hr].height = 30
    return ws


# ---- 1. Start here ---------------------------------------------------------
steps = [
    ["1", "DataForSEO", "Send the DataForSEO login and password (one account, all clients).",
     "You, once",
     "The one SEO data source never supplied. We fund it ourselves today - every grid "
     "probe, keyword volume and backlink figure is billed to us."],
    ["2", "A working email account",
     "Send IMAP host, username and password for a mailbox that actually authenticates. "
     "Ideally a catch-all domain mailbox, or Gmail with IMAP plus an app password.",
     "You, once",
     "The account sent recently does not work. Until it is replaced, no directory or "
     "Web 2.0 account can be created or confirmed."],
    ["3", "Platform accounts",
     f"Create accounts on the platforms you want used: {T['cit_signups']} citation domains "
     f"(tab 'Citation accounts') and {T['web2_accounts']} Web 2.0 platforms "
     "(tab 'Web 2.0 accounts').",
     "You, per platform",
     "The extension publishes from a logged-in browser session. No account, no placement."],
    ["4", "Each client's business record",
     "Name, address, phone, hours, categories, description, logo and photos.",
     "You, per client",
     "Every citation submission is filled from this one record, so details never drift "
     "between directories."],
    ["-", "Nothing else", "See tab 'Not needed' for the 10 items revision 1 asked for "
     "that are not part of this platform.", "Nobody", "Removed from the ask."],
]
sheet(
    "Start here",
    "AIOS - Access & Accounts (revision 2)",
    "Replaces revision 1. Platform counts read from the running registry on 12 September "
    "2026. Everything publishes through the browser extension; direct APIs are a fallback.",
    ["#", "What", "What to do", "Who", "Why it matters"],
    steps, [5, 26, 72, 20, 60],
)

# ---- 2. Core APIs ----------------------------------------------------------
rows = []
for n in D.NEEDED:
    rows.append([n["name"], n["area"], "NEEDED", n["powers"], n["key"],
                 "Agency (one account, all clients)", n["where"], n["cost"], n["why"], ""])
for nm, cat, powers, key in D.INSTALLED:
    rows.append([nm, cat, "IN PLACE", powers, key,
                 "Agency (one account, all clients)", "Already held by us", "Paid by us",
                 "", "n/a"])
sheet(
    "Core APIs", "Core platform APIs",
    f"{T['needed_keys']} needed from you, {T['installed_keys']} already installed. "
    "All model calls go through AgentRouter rather than direct to Anthropic.",
    ["Service", "Area", "Status", "What it powers", "Key / fields", "Scope",
     "Where to get it", "Cost", "Why we are asking", "Provided? (tick)"],
    rows, [26, 24, 11, 56, 44, 30, 34, 18, 62, 15],
)

# ---- 3. Not needed ---------------------------------------------------------
sheet(
    "Not needed", "Removed from the ask",
    "Revision 1 asked for all of these. None exist in this platform, or none are needed. "
    "Listed rather than quietly dropped.",
    ["Item", "Why it is not needed"],
    [[n, r] for n, r in D.REMOVED], [34, 96],
)

# ---- 4. Web 2.0 accounts ---------------------------------------------------
w2 = []
for r in D.W2_USE:
    w2.append([
        r["name"],
        "Account needed" if r["account"] == "Account needed" else "No account (anonymous)",
        r["authority"], "Extension (primary)",
        "yes - used when it works" if r["api_fallback"] == "yes" else "no",
        r["signup"], r["notes"][:180], "",
    ])
sheet(
    "Web 2.0 accounts", "Web 2.0 - platforms to create accounts on",
    f"{T['web2_use']} platforms in use; {T['web2_accounts']} need an account. Every one of "
    f"them publishes through the EXTENSION. {T['web2_api_fallback']} also have a direct API "
    "we use as a fallback when it happens to work - it is never the primary path.",
    ["Platform", "Account", "Authority", "How we publish", "Direct API fallback",
     "Sign up at", "Notes", "Created? (tick)"],
    w2, [30, 22, 12, 20, 22, 38, 70, 15],
)

# ---- 5. Web 2.0 excluded ---------------------------------------------------
sheet(
    "Web 2.0 excluded", "Web 2.0 - do NOT create accounts here",
    f"{T['web2_excluded']} platforms are catalogued but excluded on purpose - paste sites, "
    "bookmarking and curation services that do not earn a placement. Listed so the "
    "exclusion is visible rather than looking like an oversight.",
    ["Platform", "Authority", "Why excluded", "Sign up at (not needed)"],
    [[r["name"], r["authority"], r["notes"][:180], r["signup"]] for r in D.W2_EXCL],
    [30, 12, 80, 38],
)

# ---- 6. Citation accounts (deduped by domain) ------------------------------
cit = []
for r in D.CIT_SIGNUPS:
    markets = ", ".join(dict.fromkeys(r["also"]))
    cit.append([
        r["name"], r["host"], markets, r["tier"],
        "open signup" if "open" in r["account"] else "application reviewed",
        "yes - operator solves it" if r["captcha"] else "no",
        r["price"], len(r["also"]), r["verticals"], "",
    ])
sheet(
    "Citation accounts", "Citations - domains to create accounts on",
    f"{T['cit_total']} directory listings in the registry resolve to {T['cit_signups']} "
    "sign-ups, because per-market listings on the same domain share one login. Sign up with "
    "an alias of the shared mailbox. The operator submits every listing through the "
    "extension - there is no submission API and no bot.",
    ["Directory", "Domain", "Markets covered", "Authority tier", "Signup",
     "CAPTCHA", "Price", "Listings this account covers", "Verticals", "Created? (tick)"],
    cit, [34, 30, 20, 14, 22, 22, 26, 16, 26, 15],
)

# ---- 7. All citation listings ---------------------------------------------
allc = []
for r in D.CIT:
    allc.append([
        r["name"], r["url"], r["market"], r["tier"], r["account"],
        "yes" if r["captcha"] else "no", "manual only" if r["manual"] else "extension",
        r["price"], r["notes"][:160],
    ])
sheet(
    "All citation listings", "Citations - every listing in the registry",
    f"All {T['cit_total']} rows, including the {T['cit_aggregator']} aggregator entries that "
    "need nothing from you. Use the 'Citation accounts' tab for the actual sign-up list.",
    ["Directory", "URL", "Market", "Authority tier", "Account", "CAPTCHA", "How we submit",
     "Price", "Notes"],
    allc, [34, 40, 10, 14, 34, 11, 16, 26, 64],
)

# ---- 8. How the extension works --------------------------------------------
flow = [
    ["1", "Content is created in AIOS",
     "Article, title, body, images - or the client's business record for a citation. "
     "Approved before anything moves.", "Platform"],
    ["2", "It is carried into the extension",
     "Every block of the approved content is copied into the extension side panel, ready "
     "to place.", "Extension"],
    ["3", "The operator opens the platform",
     "Any citation directory or any Web 2.0 site, in their own logged-in browser. This is "
     "why the accounts are needed.", "Operator"],
    ["4", "Fields fill themselves",
     "A known field spec first, then the page's own field attributes, then AI matching for "
     "whatever is left. Values are read back to confirm they landed.", "Extension"],
    ["5", "A person presses submit",
     "Never automated. Directories forbid automated submission, and a bot that solves a "
     "CAPTCHA breaks their terms.", "Operator"],
    ["6", "The result is verified",
     "The public URL is fetched and the link checked before the placement counts.",
     "Platform"],
]
sheet(
    "How the extension works", "The publishing flow - citations and Web 2.0",
    "One mechanism for both lanes. Direct APIs exist for some Web 2.0 platforms and are "
    "used when they work, but they never reach a 100% success rate, so the extension is "
    "the primary path for everything.",
    ["Step", "What happens", "Detail", "Who does it"],
    flow, [6, 34, 88, 16],
)

wb.save(OUT)
print("WROTE", OUT)
print("size", OUT.stat().st_size, "bytes")
print("sheets", wb.sheetnames)
print("core", len(rows), "| web2", len(w2), "| citation accounts", len(cit),
      "| all listings", len(allc))
